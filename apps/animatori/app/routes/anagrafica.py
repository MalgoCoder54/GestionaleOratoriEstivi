import io
import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path

from flask import Blueprint, current_app, g, jsonify, render_template, request, send_file

from ..config_manager import get_current_event_id, load_config
from ..models import Animatore, ContributoAnimatore, DisponibilitaAnimatore, db
from ..request_context import get_animatore_or_404

# Import delle funzioni del report standalone (genera_report_animatori.py nella root webapp)
_REPO_ROOT = Path(__file__).resolve().parent.parent.parent
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))
import genera_report_animatori as _rpt  # noqa: E402
from openpyxl import Workbook  # noqa: E402

bp = Blueprint("anagrafica", __name__)

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_NO_VALUES = {"", "no", "n/a", "na", "nessuno", "nessuna"}


def _actor_name():
    auth_user = getattr(g, "auth_user", None) or {}
    return auth_user.get("name") or "App"


def _supports_in_attesa_firma():
    row = db.session.execute(db.text("""
        SELECT TOP (1)
            CASE
                WHEN cc.[definition] LIKE '%IN_ATTESA_FIRMA%' THEN CAST(1 AS BIT)
                ELSE CAST(0 AS BIT)
            END AS [SupportsAttesaFirma]
        FROM sys.check_constraints cc
        INNER JOIN sys.tables t
            ON t.[object_id] = cc.[parent_object_id]
        INNER JOIN sys.schemas s
            ON s.[schema_id] = t.[schema_id]
        WHERE s.[name] = 'animatori'
          AND t.[name] = 'animatori'
          AND cc.[name] = 'CK_animatori_StatoOperativo'
    """)).first()
    return bool(row and row.SupportsAttesaFirma)


def _coerce_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "si", "sì", "yes", "on"}
    return bool(value)


def _parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).split("T")[0])


def _clean_text(value, uppercase=False, null_no=False):
    if value is None:
        return None
    text = str(value).strip()
    if null_no and text.lower() in _NO_VALUES:
        return None
    if not text:
        return None
    return text.upper() if uppercase else text


def _validate_payload(data):
    errors = []
    if not isinstance(data, dict):
        return ["Payload JSON non valido."]
    if not _clean_text(data.get("Nome")):
        errors.append("Nome e' obbligatorio.")
    if not _clean_text(data.get("Cognome")):
        errors.append("Cognome e' obbligatorio.")
    for field in ("EmailModuli", "MailMamma", "MailPapa"):
        value = _clean_text(data.get(field), null_no=True)
        if value and not _EMAIL_RE.match(value):
            errors.append(f"{field} non e' un indirizzo email valido.")
    try:
        _parse_date(data.get("DataNascita"))
    except ValueError:
        errors.append("DataNascita non e' una data valida.")
    return errors


def _apply_animatore_fields(animatore, data):
    text_fields = {
        "Cellulare", "EmailModuli", "TagliaMaglietta", "TagliaPantaloncini",
        "AllergieIntolleranze", "TerapieNote", "NomeMamma", "CognomeMamma",
        "MailMamma", "CellulareMamma", "NomePapa", "CognomePapa",
        "MailPapa", "CellularePapa", "StatoDocumenti", "StatoOperativo",
        "NoteSegreteria",
    }
    for field in ("Nome", "Cognome"):
        if field in data:
            setattr(animatore, field, _clean_text(data.get(field)) or "")
    if "CodiceFiscale" in data:
        animatore.CodiceFiscale = _clean_text(data.get("CodiceFiscale"), uppercase=True)
    if "DataNascita" in data:
        animatore.DataNascita = _parse_date(data.get("DataNascita"))
    for field in text_fields:
        if field in data:
            setattr(animatore, field, _clean_text(data.get(field), null_no=field.startswith(("Nome", "Cognome", "Mail", "Cellulare"))))
    for field in ("Navetta", "Maggiorenne", "IscrizioneValidata", "MagliettaConsegnata"):
        if field in data:
            setattr(animatore, field, _coerce_bool(data.get(field)))
    if "DataValidazione" in data:
        animatore.DataValidazione = _parse_date(data.get("DataValidazione"))
    animatore.ModificatoDa = _actor_name()


def _create_default_children(animatore, config):
    contributo = ContributoAnimatore(
        ID_Animatore=animatore.ID,
        ID_Evento=animatore.ID_Evento,
        ImportoContributo=config["importi_default"]["contributo"],
        TotaleDovuto=config["importi_default"]["contributo"],
        ModificatoDa=_actor_name(),
    )
    db.session.add(contributo)
    for week in range(1, config["settimane"]["numero_settimane"] + 1):
        db.session.add(DisponibilitaAnimatore(
            ID_Animatore=animatore.ID,
            ID_Evento=animatore.ID_Evento,
            NumeroSettimana=week,
            ModificatoDa=_actor_name(),
        ))


def _serialize_animatore(animatore):
    data = animatore.to_dict()
    data["settimane"] = [{
        "ID": item.ID,
        "NumeroSettimana": item.NumeroSettimana,
        "Disponibile": item.Disponibile,
        "Presente": item.Presente,
        "InGita": item.InGita,
        "InOratorio": item.InOratorio,
        "NoteTurno": item.NoteTurno,
    } for item in animatore.settimane]
    if animatore.contributo:
        data["contributo"] = {
            "Pagato": animatore.contributo.Pagato,
            "TotaleDovuto": float(animatore.contributo.TotaleDovuto or 0),
            "ContabileRicevuta": animatore.contributo.ContabileRicevuta,
        }
    return data


@bp.route("/")
def home():
    evento = get_current_event_id()
    config = load_config()
    animatori = (
        Animatore.query.filter_by(ID_Evento=evento)
        .order_by(Animatore.Cognome, Animatore.Nome)
        .all()
    )
    return render_template("home.html", animatori=animatori, config=config)


@bp.route("/api/animatori")
def api_lista_animatori():
    evento = get_current_event_id()
    q = request.args.get("q", "").strip()
    query = Animatore.query.filter_by(ID_Evento=evento)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(Animatore.Nome.ilike(like), Animatore.Cognome.ilike(like), Animatore.CodiceFiscale.ilike(like)))
    rows = query.order_by(Animatore.Cognome, Animatore.Nome).limit(500).all()
    return jsonify([{
        "ID": item.ID,
        "Nome": item.Nome,
        "Cognome": item.Cognome,
        "Maggiorenne": item.Maggiorenne,
        "Cellulare": item.Cellulare,
        "StatoDocumenti": item.StatoDocumenti,
        "Pagato": bool(item.contributo and item.contributo.Pagato),
    } for item in rows])


@bp.route("/api/animatori/<int:id>")
def api_dettaglio_animatore(id):
    return jsonify(_serialize_animatore(get_animatore_or_404(id)))


@bp.route("/api/animatori", methods=["POST"])
def api_crea_animatore():
    data = request.get_json(silent=True)
    errors = _validate_payload(data)
    if errors:
        return jsonify({"errors": errors}), 400
    try:
        config = load_config()
        animatore = Animatore(ID_Evento=get_current_event_id())
        _apply_animatore_fields(animatore, data)
        animatore.StatoDocumenti = "INVIATI"
        animatore.StatoOperativo = "IN_ATTESA_FIRMA" if _supports_in_attesa_firma() else "IMPORTATO"
        db.session.add(animatore)
        db.session.flush()
        _create_default_children(animatore, config)
        db.session.commit()
        return jsonify({"id": animatore.ID}), 201
    except Exception:
        current_app.logger.exception("Errore creazione animatore")
        db.session.rollback()
        return jsonify({"error": "Errore creazione animatore."}), 500


@bp.route("/api/animatori/<int:id>", methods=["PUT"])
def api_modifica_animatore(id):
    animatore = get_animatore_or_404(id)
    data = request.get_json(silent=True)
    errors = _validate_payload(data)
    if errors:
        return jsonify({"errors": errors}), 400
    _apply_animatore_fields(animatore, data)
    animatore.DataModifica = datetime.now(timezone.utc)
    db.session.commit()
    return jsonify({"ok": True})


@bp.route("/api/export/animatori")
def export_animatori_excel():
    """Genera il report Excel completo degli animatori (4 fogli)."""
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    raw_conn = db.engine.raw_connection()
    try:
        cursor = raw_conn.cursor()
        _rpt.foglio_anagrafica(wb, cursor)
        _rpt.foglio_magliette(wb, cursor)
        _rpt.foglio_allergie(wb, cursor)
        _rpt.foglio_navetta(wb, cursor)
        cursor.close()
    finally:
        raw_conn.close()

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="report_animatori.xlsx",
    )


@bp.route("/api/animatori/<int:id>", methods=["DELETE"])
def api_elimina_animatore(id):
    animatore = get_animatore_or_404(id)
    try:
        db.session.execute(
            db.text("UPDATE [animatori].[import_animatori_forms_log] SET [ID_Animatore] = NULL WHERE [ID_Animatore] = :id"),
            {"id": id},
        )
        db.session.delete(animatore)
        db.session.commit()
        return jsonify({"ok": True})
    except Exception:
        current_app.logger.exception("Errore eliminazione animatore")
        db.session.rollback()
        return jsonify({"error": "Errore eliminazione animatore."}), 500
